"""Per-enquiry export (CSV / .xlsx) for enquiry modules.

Unlike the Team Daily Tracker export (``tracker_export.py``, per-member daily
counts), this writes one row per enquiry so users can analyse the raw data.
The file is built entirely server-side from the same filtered queryset as the
list endpoint (no pagination), in either format:

    ?file_format=csv   (default) UTF-8 with BOM so Excel opens it cleanly
    ?file_format=xlsx  typed cells: numbers, real dates / datetimes

``file_format`` is used instead of ``format`` because DRF reserves ``?format=``
for renderer negotiation (an unknown value 404s).

A module plugs in by declaring a column spec: ``(header, getter, kind)`` where
``getter(entry)`` returns a native value (str / int / Decimal / date / local
naive datetime / None) and ``kind`` is one of ``text``, ``int``, ``number``,
``date``, ``datetime``. Each writer formats values for its own format.
"""
from __future__ import annotations

import csv
from io import BytesIO, StringIO

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from rest_framework.exceptions import ParseError

from .models import MarineNewEntry
from .tracker_export import BORDER, FILL_HEADER

FILE_FORMATS = ('csv', 'xlsx')
XLSX_CONTENT_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

# Leading characters a spreadsheet may interpret as a formula (CSV injection).
FORMULA_PREFIXES = ('=', '+', '-', '@', '\t', '\r')

XLSX_NUMBER_FORMATS = {
    'int': '0',
    'number': '#,##0.00',
    'date': 'yyyy-mm-dd',
    'datetime': 'yyyy-mm-dd hh:mm',
}


def _local(dt):
    """Aware datetime -> naive datetime in the business timezone (Asia/Dubai)."""
    if dt is None:
        return None
    return timezone.localtime(dt).replace(tzinfo=None)


def _csv_value(value, kind):
    if value is None:
        return ''
    if kind == 'date':
        return value.isoformat()
    if kind == 'datetime':
        return value.strftime('%Y-%m-%d %H:%M')
    if kind == 'text':
        value = str(value)
        if value.startswith(FORMULA_PREFIXES):
            return "'" + value
    return value


def csv_response(filename, columns, rows):
    buf = StringIO()
    buf.write('﻿')  # BOM: Excel otherwise mis-decodes UTF-8
    writer = csv.writer(buf)
    writer.writerow([header for header, _, _ in columns])
    for entry in rows:
        writer.writerow([_csv_value(get(entry), kind) for _, get, kind in columns])

    response = HttpResponse(buf.getvalue(), content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


def xlsx_response(filename, sheet_title, columns, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]

    for col, (header, _, _) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = FILL_HEADER
        cell.border = BORDER

    row_count = 0
    for r, entry in enumerate(rows, start=2):
        row_count += 1
        for col, (_, get, kind) in enumerate(columns, start=1):
            value = get(entry)
            if value is None or value == '':
                continue
            cell = ws.cell(row=r, column=col)
            if kind == 'text':
                # Force the string type so "=..." stays literal text instead
                # of becoming a live formula (openpyxl infers 'f' from "=").
                # Control characters (e.g. a vertical tab pasted from Word)
                # are illegal in XML and would make openpyxl raise.
                cell.value = ILLEGAL_CHARACTERS_RE.sub('', str(value))
                cell.data_type = 's'
            else:
                cell.value = float(value) if kind == 'number' else value
                cell.number_format = XLSX_NUMBER_FORMATS[kind]

    for col, (header, _, kind) in enumerate(columns, start=1):
        width = 18 if kind == 'datetime' else max(12, min(len(header) + 4, 40))
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(columns))}{row_count + 1}'

    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type=XLSX_CONTENT_TYPE)
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


def export_response(request, basename, sheet_title, columns, rows):
    """Dispatch on ``?file_format=`` (csv default). 400 on anything else."""
    file_format = (request.query_params.get('file_format') or 'csv').lower()
    if file_format not in FILE_FORMATS:
        raise ParseError('file_format must be csv or xlsx')
    filename = f'{basename}_{timezone.localdate().isoformat()}.{file_format}'
    if file_format == 'xlsx':
        return xlsx_response(filename, sheet_title, columns, rows)
    return csv_response(filename, columns, rows)


# ---- Marine New ----------------------------------------------------------

# Mirrors STATUS_CONFIG['marine-new'] in MarineNewEnquiryPage.tsx (the UI
# shows "Won" where the model's choice label says "Converted").
MARINE_NEW_STATUS_LABELS = {
    MarineNewEntry.STATUS_NEW: 'New Enquiry',
    MarineNewEntry.STATUS_SHARED_WITH_CLIENT: 'Shared With Client',
    MarineNewEntry.STATUS_CONVERTED: 'Won',
    MarineNewEntry.STATUS_REJECTED: 'Rejected',
    MarineNewEntry.STATUS_LOST: 'Lost',
}


def _marine_status(entry):
    return MARINE_NEW_STATUS_LABELS.get(entry.status) or entry.get_status_display()


def _compared_insurers(entry):
    # Iterates the prefetched M2M (no per-row query).
    return sorted(entry.compared_insurance_companies.all(), key=lambda c: c.name)


def _insurer_refs(entry):
    """Same rule as MarineNewEntrySerializer.get_insurer_references_display:
    skip ids no longer among the compared insurers, order by insurer name."""
    refs = entry.insurer_references or {}
    pairs = [
        (company.name, refs[str(company.id)])
        for company in _compared_insurers(entry)
        if str(company.id) in refs
    ]
    return '; '.join(f'{name}: {ref}' for name, ref in pairs)


def _name(user):
    return user.get_full_name() if user else None


MARINE_NEW_COLUMNS = [
    ('ID', lambda e: e.pib_id, 'text'),
    # Linked client's name, else the legacy typed name.
    ('Client Name', lambda e: e.client_display, 'text'),
    ('Status', _marine_status, 'text'),
    ('Insurer Ref.', _insurer_refs, 'text'),
    ('Potential Premium', lambda e: e.potential_premium, 'number'),
    ('Converted Premium', lambda e: e.converted_premium, 'number'),
    ('Revisions', lambda e: e.revisions, 'int'),
    ('No. of Quotes Compared', lambda e: e.quotes_compared, 'int'),
    ('TAT (minutes)', lambda e: e.tat_minutes, 'number'),
    ('Accuracy (%)', lambda e: e.accuracy, 'number'),
    ('Added by', lambda e: _name(e.added_by), 'text'),
    ('Agent Name', lambda e: _name(e.agent), 'text'),
    ('Class of Insurance', lambda e: e.class_of_insurance.name if e.class_of_insurance else None, 'text'),
    ('Compared Insurers', lambda e: '; '.join(c.name for c in _compared_insurers(e)), 'text'),
    ('Converted Insurer', lambda e: e.converted_insurer.name if e.converted_insurer else None, 'text'),
    ('Policy/Certificate No.', lambda e: e.policy_number, 'text'),
    ('Added on', lambda e: _local(e.added_at), 'datetime'),
    ('Date', lambda e: e.date, 'date'),
    ('Shared With Client At', lambda e: _local(e.shared_with_client_at), 'datetime'),
    ('Closed At', lambda e: _local(e.status_changed_at), 'datetime'),
    ('Last Updated', lambda e: _local(e.updated_at), 'datetime'),
    ('Voided', lambda e: 'Yes' if e.is_voided else 'No', 'text'),
    ('Voided At', lambda e: _local(e.voided_at), 'datetime'),
    ('Voided By', lambda e: _name(e.voided_by), 'text'),
    ('Void Reason', lambda e: e.void_reason, 'text'),
]
